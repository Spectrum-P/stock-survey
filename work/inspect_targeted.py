from pathlib import Path
from openpyxl import load_workbook

path = Path('/Users/mac/Downloads/Stock_Condition_Survey_System (3).xlsx')
wb = load_workbook(path, data_only=False, read_only=False)

for sheet_name in ['📋 Survey Form', '📈 Dashboard', '📌 Rating Reference', '📖 Life Expectancy']:
    ws = wb[sheet_name]
    print(f'\n===== {sheet_name} {ws.calculate_dimension()} =====')
    for row in ws.iter_rows():
        cells = []
        for cell in row:
            if cell.value not in (None, ''):
                text = str(cell.value).replace('\n', ' / ')
                cells.append(f'{cell.coordinate}={text[:260]}')
        if cells:
            print(' | '.join(cells))
    dvs = [f'{dv.sqref} :: {dv.type} :: {dv.formula1}' for dv in ws.data_validations.dataValidation]
    if dvs:
        print('VALIDATIONS')
        print('\n'.join(dvs))
