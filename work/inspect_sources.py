from pathlib import Path
import sys
from docx import Document
from openpyxl import load_workbook

XLSX_FILES = [
    Path('/Users/mac/Downloads/Stock_Condition_Survey_Prototype.xlsx'),
    Path('/Users/mac/Downloads/Stock_Condition_Survey_System.xlsx'),
    Path('/Users/mac/Downloads/Stock_Condition_Survey_System (1).xlsx'),
    Path('/Users/mac/Downloads/Stock_Condition_Survey_System (3).xlsx'),
]

def clean(value):
    if value is None:
        return ''
    text = str(value).replace('\n', ' / ').strip()
    return text if len(text) <= 180 else text[:177] + '...'

def meaningful_rows(ws, max_rows=80):
    found = []
    for row in ws.iter_rows():
        values = [clean(c.value) for c in row]
        while values and not values[-1]:
            values.pop()
        if any(values):
            found.append(values)
        if len(found) >= max_rows:
            break
    return found

def workbook_summary(path):
    wb = load_workbook(path, data_only=False, read_only=False)
    print(f'\n===== WORKBOOK: {path.name} =====')
    print('Sheets:', ', '.join(wb.sheetnames))
    for ws in wb.worksheets:
        rows = meaningful_rows(ws, 15 if ws.title.startswith('_') else 80)
        formulas = []
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith('='):
                    formulas.append(f'{cell.coordinate}:{clean(cell.value)}')
        validations = [f'{dv.sqref} {dv.type} {clean(dv.formula1)}' for dv in ws.data_validations.dataValidation]
        print(f'\n-- {ws.title} | declared {ws.calculate_dimension()} | meaningful rows shown {len(rows)} | formulas {len(formulas)} | validations {len(validations)}')
        for idx, vals in enumerate(rows, 1):
            print(f'R{idx}: ' + ' || '.join(vals))
        if formulas:
            print('Formula samples:', ' | '.join(formulas[:20]))
        if validations:
            print('Validations:', ' | '.join(validations[:20]))

def docx_summary(path):
    doc = Document(path)
    print(f'===== DOCX: {path.name} =====')
    for p in doc.paragraphs:
        if p.text.strip():
            print(clean(p.text))
    for i, table in enumerate(doc.tables, 1):
        print(f'-- TABLE {i} --')
        for row in table.rows:
            print(' || '.join(clean(c.text) for c in row.cells))

def main():
    if len(sys.argv) == 1 or sys.argv[1] == 'docx':
        docx_summary(Path('/Users/mac/Downloads/STOCK CONDITION TEMPLATE.docx'))
    if len(sys.argv) == 1:
        selected = XLSX_FILES
    elif sys.argv[1].isdigit():
        selected = [XLSX_FILES[int(sys.argv[1])]]
    else:
        selected = []
    for path in selected:
        workbook_summary(path)

if __name__ == '__main__':
    main()
